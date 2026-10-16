"""Build a niche's downloadable Excel template and read a filled one back.

Workbook sheets:
    Content    one row per video; the only sheet read on upload
    Guide      what to write in every column, with limits
    Example    one complete, valid video
    AI Prompt  instructions to paste into an AI tool together with the file
"""

import datetime as dt
import io
import uuid
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Union

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from content_templates.specs import VIDEO_FIELDS, Field, NicheSpec

CONTENT_SHEET = "Content"
DATA_ROWS = 500  # rows that carry dropdowns for choice columns
PUBLISH_FORMATS = ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M", "%d-%m-%Y %H:%M", "%d/%m/%Y %H:%M")

GROUP_FILLS = {"YouTube": "1F4E79", "Intro slide": "7030A0"}
ITEM_FILLS = ("2E75B6", "548235")  # alternate per item so groups are easy to tell apart


# ---------- download ----------

def build_workbook(spec: NicheSpec) -> Workbook:
    wb = Workbook()
    content = wb.active
    content.title = CONTENT_SHEET
    _write_header(content, spec)
    for name, _, f in spec.columns():
        if f.choices:
            col = _column_index(spec, name)
            dv = DataValidation(type="list", formula1='"' + ",".join(f.choices) + '"', allow_blank=not f.required)
            dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{DATA_ROWS + 1}")
            content.add_data_validation(dv)

    _write_guide(wb.create_sheet("Guide"), spec)

    example = wb.create_sheet("Example")
    _write_header(example, spec)
    for col, (name, _, _) in enumerate(spec.columns(), start=1):
        value = spec.example.get(name, "")
        example.cell(row=2, column=col, value=value).alignment = Alignment(wrap_text=True, vertical="top")

    prompt = wb.create_sheet("AI Prompt")
    prompt.column_dimensions["A"].width = 120
    cell = prompt.cell(row=1, column=1, value=build_prompt(spec))
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    prompt.row_dimensions[1].height = 409  # Excel's maximum row height
    return wb


def workbook_bytes(spec: NicheSpec) -> bytes:
    buffer = io.BytesIO()
    build_workbook(spec).save(buffer)
    return buffer.getvalue()


def build_prompt(spec: NicheSpec) -> str:
    items = spec.items
    count = f"exactly {items.min_count}" if items.min_count == items.max_count else f"{items.min_count} to {items.max_count}"
    header = "\t".join(name for name, _, _ in spec.columns())
    rules = "\n".join(f"- {name}: {_rule_text(f)}" for name, _, f in spec.columns())
    return f"""You are writing YouTube Shorts content for the channel "{spec.name}".

Create [NUMBER OF VIDEOS] videos about: [TOPIC OR THEME].

Audience and tone: {spec.audience_and_tone}
Language: {spec.language}

Each video is one row. It has an intro slide and {count} {items.label}s. Fill {items.prefix}1 onward in order and leave the columns of unused {items.label}s empty.

How to answer:
- If you can edit the attached Excel file, fill its "{CONTENT_SHEET}" sheet from row 2, one row per video, and keep the header row exactly as it is. The "Example" sheet shows one finished video.
- Otherwise reply with only a tab-separated table inside one code block: the header line below first, then one line per video. It will be pasted into the {CONTENT_SHEET} sheet at cell A1.

Header (tab-separated, this exact order):
{header}

Columns:
{rules}

Rules:
- Stay within every character limit; spaces count.
- No line breaks inside a cell. In youtube_description, write \\n where a new line should start.
- No markdown, no quotation marks around values, no extra columns.
- Leave video_id, publish_at and every image column empty.
- Every fact must be accurate. Do not invent statistics, quotes, dates or sources.
- Every video covers a different topic; never repeat a title.
"""


def _rule_text(f: Field) -> str:
    parts = ["required" if f.required else "optional"]
    if f.max_len:
        parts.append(f"max {f.max_len} characters")
    if f.choices:
        parts.append("one of " + ", ".join(f.choices))
    return f"{'; '.join(parts)}. {f.description}"


def _write_header(sheet, spec: NicheSpec) -> None:
    item_number = 0
    previous_group = None
    for col, (name, group, f) in enumerate(spec.columns(), start=1):
        if group not in GROUP_FILLS and group != previous_group:
            item_number += 1
        previous_group = group
        color = GROUP_FILLS.get(group, ITEM_FILLS[item_number % 2])
        cell = sheet.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(vertical="center")
        cell.comment = Comment(f"{group}\n{_rule_text(f)}", "Content template", width=320, height=140)
        sheet.column_dimensions[get_column_letter(col)].width = 18 if (f.max_len or 0) <= 40 else 40
    sheet.freeze_panes = "B2"


def _write_guide(sheet, spec: NicheSpec) -> None:
    rows = [
        (f"{spec.name} — content template", None, None, None, None),
        (f"Template: {spec.template_id} ({spec.aspect_ratio}). Language: {spec.language}.", None, None, None, None),
        (f"One row per video on the {CONTENT_SHEET} sheet. Each video has {spec.items.min_count}-{spec.items.max_count} {spec.items.label}s; leave unused ones empty.", None, None, None, None),
        (None, None, None, None, None),
        ("Column", "Part", "Required", "Max characters", "What to write"),
    ]
    rows += [
        (name, group, "yes" if f.required else "optional", f.max_len or "", f.description + (f" Allowed: {', '.join(f.choices)}." if f.choices else ""))
        for name, group, f in spec.columns()
    ]
    for r, values in enumerate(rows, start=1):
        for c, value in enumerate(values, start=1):
            if value is not None:
                cell = sheet.cell(row=r, column=c, value=value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(row=1, column=1).font = Font(bold=True, size=14)
    for c in range(1, 6):
        sheet.cell(row=5, column=c).font = Font(bold=True)
    for letter, width in zip("ABCDE", (24, 14, 10, 15, 90)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A6"


def _column_index(spec: NicheSpec, name: str) -> int:
    return [n for n, _, _ in spec.columns()].index(name) + 1


# ---------- upload ----------

@dataclass
class ImportResult:
    videos: List[Dict[str, Any]] = field(default_factory=list)
    errors: List[Dict[str, Any]] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def read_workbook(source: Union[str, bytes, io.BytesIO], spec: NicheSpec) -> ImportResult:
    """Validate a filled template and convert each row to a video's content."""
    if isinstance(source, bytes):
        source = io.BytesIO(source)
    result = ImportResult()
    try:
        wb = load_workbook(source, data_only=True)
    except Exception as e:
        result.errors.append(_error(None, None, f"not a readable .xlsx file ({e})"))
        return result
    sheet = wb[CONTENT_SHEET] if CONTENT_SHEET in wb.sheetnames else wb.worksheets[0]

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        result.errors.append(_error(None, None, "the sheet is empty"))
        return result
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    known = {name: f for name, _, f in spec.columns()}
    for h in headers:
        if h and h not in known:
            result.errors.append(_error(1, h, "unknown column; check the spelling against the template"))
    must_have = [name for name, _, f in spec.columns() if not name.startswith(spec.items.prefix)]
    must_have += [f"{spec.items.prefix}{n}_{f.key}" for n in range(1, spec.items.min_count + 1) for f in spec.items.fields]
    for name in must_have:
        if name not in headers:
            result.errors.append(_error(1, name, "column is missing"))
    if result.errors:
        return result

    seen_ids = set()
    for row_number, values in enumerate(rows[1:], start=2):
        cells = {h: values[i] if i < len(values) else None for i, h in enumerate(headers) if h}
        if all(_is_blank(v) for v in cells.values()):
            continue
        errors_before = len(result.errors)
        video = _read_row(row_number, cells, spec, result.errors)
        if video["video_id"] in seen_ids:
            result.errors.append(_error(row_number, "video_id", f"duplicate video_id '{video['video_id']}'"))
        seen_ids.add(video["video_id"])
        if len(result.errors) == errors_before:
            result.videos.append(video)
    if not result.videos and not result.errors:
        result.errors.append(_error(None, None, f"no videos found; fill the {CONTENT_SHEET} sheet from row 2"))
    return result


def _read_row(row: int, cells: Dict[str, Any], spec: NicheSpec, errors: List[Dict[str, Any]]) -> Dict[str, Any]:
    def value(name: str, f: Field) -> Any:
        return _check(row, name, f, cells.get(name), errors)

    video = {name: value(name, f) for name, f in ((f.key, f) for f in VIDEO_FIELDS)}
    intro = {f.key: value(f"intro_{f.key}", f) for f in spec.intro_fields}

    slides, gap_at = [], None
    for n in range(1, spec.items.max_count + 1):
        names = {f.key: f"{spec.items.prefix}{n}_{f.key}" for f in spec.items.fields}
        present = any(not _is_blank(cells.get(c)) for c in names.values())
        if not present:
            gap_at = gap_at or n
            continue
        if gap_at:
            errors.append(_error(row, f"{spec.items.prefix}{n}", f"{spec.items.label} {n} is filled but {spec.items.label} {gap_at} is empty; fill them in order"))
            gap_at = None
        slides.append(_to_slide({key: value(col, f) for (key, col), f in zip(names.items(), spec.items.fields)}, len(slides)))
    if len(slides) < spec.items.min_count:
        errors.append(_error(row, None, f"{len(slides)} {spec.items.label}s; this template needs at least {spec.items.min_count}"))

    return {
        "row": row,
        "video_id": video["video_id"] or f"{spec.niche}-{uuid.uuid4().hex[:8]}",
        "publish_at": video["publish_at"] or None,
        "youtube": {
            "title": video["youtube_title"],
            "description": video["youtube_description"],
            "tags": [t.strip() for t in (video["youtube_tags"] or "").split(",") if t.strip()],
        },
        "content": {"intro": intro, "slides": slides},
    }


def _to_slide(values: Dict[str, Any], index: int) -> Dict[str, Any]:
    """Rename spreadsheet fields to the keys the template's data uses."""
    slide = dict(values)
    if "option_a" in slide:
        slide["options"] = [slide.pop(k) for k in ("option_a", "option_b", "option_c", "option_d")]
    if "text_position" in slide:
        slide["textPosition"] = slide.pop("text_position") or ("left" if index % 2 == 0 else "right")
    if "image" in slide and not slide["image"]:
        slide["image"] = ""
    return slide


def _check(row: int, name: str, f: Field, raw: Any, errors: List[Dict[str, Any]]) -> Any:
    if name == "publish_at":
        return _publish_at(row, raw, errors)
    text = _text(raw, keep_newlines=(name == "youtube_description"))
    if not text:
        if f.required:
            errors.append(_error(row, name, "required"))
        return ""
    if f.max_len and len(text) > f.max_len:
        errors.append(_error(row, name, f"{len(text)} characters; the limit is {f.max_len}"))
    if f.kind == "answer":
        text = text.upper()
    if f.choices and text not in f.choices:
        errors.append(_error(row, name, f"'{text}' is not allowed; use one of {', '.join(f.choices)}"))
    if f.kind == "url" and not text.startswith(("http://", "https://")):
        errors.append(_error(row, name, "must be a full http(s) URL or empty"))
    if f.kind == "bars":
        return _bars(row, name, text, errors)
    return text


def _bars(row: int, name: str, text: str, errors: List[Dict[str, Any]]) -> List[int]:
    try:
        bars = [int(float(p)) for p in text.replace(";", ",").split(",") if p.strip()]
    except ValueError:
        bars = []
    if len(bars) != 7 or any(b < 0 or b > 100 for b in bars):
        errors.append(_error(row, name, "needs seven numbers from 0 to 100 separated by commas"))
    return bars


def _publish_at(row: int, raw: Any, errors: List[Dict[str, Any]]) -> Optional[str]:
    if _is_blank(raw):
        return None
    if isinstance(raw, dt.datetime):
        return raw.strftime("%Y-%m-%dT%H:%M")
    text = str(raw).strip()
    for fmt in PUBLISH_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).strftime("%Y-%m-%dT%H:%M")
        except ValueError:
            pass
    errors.append(_error(row, "publish_at", f"'{text}' is not a date and time; use YYYY-MM-DD HH:MM"))
    return None


def _text(raw: Any, keep_newlines: bool = False) -> str:
    if _is_blank(raw):
        return ""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    text = str(raw).replace("\r\n", "\n").replace("\\n", "\n").strip()
    if keep_newlines:
        return "\n".join(line.strip() for line in text.split("\n"))
    return " ".join(text.split())


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _error(row: Optional[int], column: Optional[str], message: str) -> Dict[str, Any]:
    return {"row": row, "column": column, "message": message}
